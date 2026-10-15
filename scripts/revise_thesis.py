from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

root=Path(__file__).resolve().parents[1]
source=root/'baocao/exports/11236160_HoangVanNam_KhoaLuan_UniMatch.docx'
output=root/'baocao/exports/11236160_HoangVanNam_KhoaLuan_UniMatch_hoan_thien.docx'
doc=Document(source)
replacements={
 'UI 3D biểu diễn đồ thị hai phía:':'Mô phỏng dùng WebGL qua Three.js 0.186.1 với camera phối cảnh, vật liệu có phản xạ và OrbitControls [15]. Mặt sàn bố trí toàn bộ sinh viên bên trái và các giảng viên cùng cột hạn mức bên phải trong một vùng hữu hạn. Cột được làm đầy theo tải thực, không dùng vị trí hình học làm khoảng cách embedding. Khi phát một quyết định được chấp nhận, hình sinh viên di chuyển tới nhóm của giảng viên; chỉ đường ghép hiện hành được chiếu sáng để tránh chồng chéo hàng trăm cạnh.',
 'Mỗi event chứa step,':'Mỗi event từ API giữ nguyên bước, ID, điểm tương thích, reward nếu có, invalid proposal, fallback và tải. Replay tái tạo tập assignment duy nhất theo sinh viên từ đầu trace đến bước đang chọn. Vì vậy, lùi bước hoặc kéo thanh thời gian không cộng dồn sai tải. Đề xuất không hợp lệ không tạo assignment; fallback được đánh dấu riêng. Tổng đã ghép và cột hạn mức là trạng thái của bước phát lại, còn bảng phía dưới trình bày kết quả hoàn chỉnh của lần chạy.',
 'Đồ thị 3D không phải giao diện duy nhất.':'Thông tin thiết yếu còn có trong lời giải thích HTML, bộ đếm và bảng kết quả. Người dùng có thể kéo để xoay, cuộn để thu phóng, chọn cột giảng viên để đọc tải, hoặc dùng phím mũi tên và dấu cộng trừ. Khi WebGL không khả dụng, matching dùng renderer Canvas dự phòng. Bố cục hữu hạn đã được kiểm tra với 154 sinh viên và 38 giảng viên; điều này không thay cho đánh giá khả dụng với người dùng ngoài nhóm phát triển.',
 'Ảnh docs/ui-reference/matching-ui-v1.png':'Ảnh docs/ui-reference/matching-ui-v1.png lưu giao diện tham chiếu trước đợt thay đổi. Bản hiện hành thay landing bằng sân khấu xanh đậm với hai cụm hồ sơ dạng khối 3D, quỹ đạo và hạt sáng nối chuyên môn. Đồ họa ở trang chủ là minh họa; cảnh matching được điều khiển bằng trace thật. Việc trang trí không làm thay đổi thuật toán, score hoặc điều kiện promotion.',
 'Trong bản cloud, frontend Vercel chỉ chứa':'Mã nguồn frontend, backend, cấu hình và lịch ping được hợp nhất trong repository công khai namhv521/UniMatch. Frontend Vercel chỉ chứa tài nguyên giao diện và thư viện Three.js tự phục vụ. Hồ sơ và checkpoint được lưu riêng trong bảng serving_artifacts của Neon dưới bundle bất biến; URI database và mật khẩu ứng dụng chỉ đặt trong môi trường Render. Không cần repository private triển khai hoặc repository uptime riêng.',
 'Các project được tạo mới trong phiên triển khai;':'Frontend được triển khai tại project Vercel unimatch, backend tại service Render unimatch-api-20261016 và database tại project Neon unimatch ở Singapore. Database được tạo riêng và không dùng database P076. Hai dịch vụ web đều đọc mã nguồn từ namhv521/UniMatch. Tên project Vercel là unimatch; domain sử dụng unimatch-neu.vercel.app vì unimatch.vercel.app đã được Vercel xác nhận thuộc project khác. URL cũ chuyển hướng tới domain mới.',
 '• Bộ kiểm tra liên quan tại source:':'• Bộ kiểm tra hiện hành: 197 test Python thành công, 1 test được bỏ qua và 1 cảnh báo thư viện; 16 test frontend thành công. Kiểm tra cú pháp JavaScript và build frontend thành công. Đây là bằng chứng chức năng ở phạm vi đã chạy, không phải kết luận về uptime hoặc bảo mật toàn diện.',
 '• Neon lưu 1 snapshot,':'• Database Neon mới được đọc lại có 1 snapshot, 215 students, 38 advisors, 215 split memberships và 92 serving_artifacts. Lịch sử simulation_runs tăng sau các request xác thực; số đếm là trạng thái tại thời điểm kiểm chứng, không phải kích thước cố định của hệ thống.',
 'Kết quả chi tiết không chứa credential:':'Kết quả kiểm chứng cuối không chứa credential được lưu tại outputs/deployment/unimatch-release.verification.json. Phân biệt bản kiểm chứng này với log từ topology triển khai trước khi hợp nhất repository.',
 'Workflow .github/workflows/keepalive.yml trong repo uptime':'Workflow .github/workflows/keepalive.yml nằm ngay trong namhv521/UniMatch và gọi endpoint health của Render mỗi năm phút. Workflow chỉ có quyền đọc nội dung repository và kiểm tra data_loaded=true; lỗi tải dữ liệu làm job thất bại. Lượt chạy thủ công là kiểm chứng đường gọi; lịch tự động phải được theo dõi bằng lịch sử Actions [13].',
 'Lịch Actions tối thiểu năm phút là lịch best effort.':'GitHub Actions có thể trì hoãn lịch và có thể vô hiệu lịch trong repository công khai không hoạt động đủ lâu [13]. Render Free có thể restart hoặc hết ngân sách sử dụng [11]. Do đó ping hỗ trợ duy trì request định kỳ nhưng không bảo đảm backend luôn thức. Timer nằm bên trong backend không giải quyết được việc ngủ vì timer cũng dừng cùng tiến trình.',
 'Trang đầu sau đăng nhập được thiết kế thành landing page':'Trang chủ sau đăng nhập sử dụng nền xanh đậm liên tục, chữ Manrope, các cụm nhân vật 3D và đường nối có hạt sáng chuyển động. Vật liệu, ánh sáng và parallax làm rõ hai nhóm chuyên môn. Nội dung và hai lối vào nghiệp vụ vẫn là HTML. Vòng vẽ dừng khi cảnh ra ngoài màn hình, tab bị ẩn hoặc người dùng yêu cầu; reduced motion được tôn trọng mặc định. Người dùng có thể chủ động bật chuyển động. Tài nguyên GPU và listener được thu hồi khi chuyển route.',
 'Mã nguồn triển khai nằm trong repo private':'Mã nguồn, lịch ping và cấu hình triển khai nằm trong https://github.com/namhv521/UniMatch. Mật khẩu ứng dụng, URI PostgreSQL, hồ sơ nguồn và checkpoint không nằm trong repository hoặc khóa luận. Quản trị viên giữ credential ở môi trường backend; bộ inference được nạp từ bundle có hash trong Neon. Khi bàn giao cần kiểm tra lại quyền truy cập, cookie, Origin và khả năng khôi phục bundle.',
 'Các artifact đối chiếu gồm outputs/benchmarks':'Các artifact đối chiếu gồm outputs/benchmarks/current-validation.json, outputs/deployment/unimatch-release.verification.json, docs/cloud_deployment_status.md và .github/workflows/keepalive.yml. Nếu đổi dữ liệu, feature, quota hoặc mapping phải tạo protocol mới thay vì ghép số liệu vào bảng cũ.',
}
for p in doc.paragraphs:
 for prefix,text in replacements.items():
  if p.text.startswith(prefix):p.text=text;break
# Historical screenshots are retained as historical evidence, not mislabeled as the WebGL release.
for p in doc.paragraphs:
 if p.text.startswith('Hình 4.') and any(w in p.text for w in ['Trang đầu','Mô phỏng Hungarian','Đọc lại snapshot','Lượt GitHub']):
  for run in p.runs:
   if run.text and 'Hình ' not in run.text:run.text=run.text.replace(' của hệ thống đã triển khai',' của bản trước khi nâng cấp WebGL')
 for run in p.runs:
  if 'Minh chứng trực quan của luồng đã kiểm chứng' in run.text:run.text='Ảnh sau ghi nhận lần triển khai trước khi hợp nhất repository và nâng cấp WebGL. Các định danh cloud trong ảnh thuộc mốc lịch sử; thông tin triển khai hiện hành lấy theo bảng 4.5.'
# Preserve automatic caption fields while improving formula names and current URLs.
formula_names={'1.1':'Mục tiêu tối đa hóa tương thích','1.2':'Ràng buộc ghép duy nhất và hạn mức','2.1':'Giá trị kỳ vọng của policy','2.2':'Mục tiêu cập nhật Q learning','2.3':'Tỷ số xác suất policy','2.4':'Mục tiêu PPO có clipping','2.5':'Trọng số TF IDF','2.6':'Độ tương tự cosine','2.7':'Kết hợp đặc trưng theo trọng số','2.8':'Định danh snapshot bằng SHA 256','2.9':'Mở rộng hạn mức thành các slot','2.10':'Quyết định tham lam hợp lệ','2.11':'Điểm tương thích trung bình','2.12':'Độ lệch tải giảng viên','2.13':'Phương sai tải giảng viên','2.14':'Tỷ số chất lượng so với Hungarian','2.15':'Biểu diễn trạng thái MDP','2.16':'Mask hành động theo hạn mức','2.17':'Hàm thưởng ghép cặp','3.1':'Hạn mức đồng đều trong thực nghiệm','4.1':'Phương sai tải tối thiểu trên validation'}
for p in doc.paragraphs:
 for run in p.runs:
  if 'https://unimatch-web-20261016.vercel.app' in run.text:run.text=run.text.replace('https://unimatch-web-20261016.vercel.app','https://unimatch-neu.vercel.app')
 m=re.match(r'Công thức (\d+\.\d+) ',p.text)
 if m and m[1] in formula_names:
  for run in p.runs:
   if 'Công thức mô hình và đánh giá' in run.text:run.text=run.text.replace('Công thức mô hình và đánh giá',formula_names[m[1]])

# Add the bootstrap design where the reader meets persistence.
anchor=next(p for p in doc.paragraphs if p.text.startswith('3.5.3 Tái lập'))
for text in [
 'Để clone mã nguồn công khai vẫn có thể triển khai mà không công bố hồ sơ, backend đọc bundle định danh bởi KLTN_ARTIFACT_BUNDLE từ bảng serving_artifacts. Bảng có khóa ghép bundle–path, SHA-256 và payload bytea. Bộ hiện hành gồm 92 file với tổng 6.612.229 byte; chỉ chứa các thư mục curated, features, checkpoint và benchmark cần phục vụ, không chứa thư mục PDF nguồn.',
 'Trước khi ghi file, chương trình kiểm đường dẫn theo allowlist, từ chối đường dẫn tuyệt đối hoặc vượt root và xác minh toàn bộ hash của bundle. Sau đó từng file được thay thế nguyên tử, MatchingService tiếp tục kiểm shape, split và mapping. Bundle thiếu hoặc hỏng khiến khởi động thất bại; không tự chấp nhận dữ liệu khác. Nhờ vậy, restart Render có thể dựng lại inference trên filesystem tạm, còn dữ liệu ứng dụng vẫn lưu ở PostgreSQL.'
]:anchor.insert_paragraph_before(text)
# Replace historical UI captures with fresh Chrome evidence when available.
for caption_key,filename in [('Hình 4.2','unimatch-live-home.jpg'),('Hình 4.3','unimatch-live-matching.jpg'),('Hình 4.5','unimatch-live-ping.jpg')]:
 paragraphs=doc.paragraphs
 for index,p in enumerate(paragraphs):
  if p.style.name.startswith('Caption') and p.text.startswith(caption_key):
   image_path=root/'outputs/deployment'/filename
   if image_path.exists():
    image_p=paragraphs[index-1]
    if image_p._p.findall('.//'+qn('w:drawing')):
     image_p.clear();image_p.add_run().add_picture(str(image_path),width=Cm(14.8))
     if index>=2 and paragraphs[index-2].text.startswith('Ảnh sau ghi nhận'):
      paragraphs[index-2].text='Ảnh chụp Chrome trên topology hiện hành sau khi hợp nhất UniMatch. Các con số hiển thị là trạng thái tại thời điểm kiểm chứng và không thay thế benchmark.'
     for run in p.runs:
      if ' của bản trước khi nâng cấp WebGL' in run.text:run.text=run.text.replace(' của bản trước khi nâng cấp WebGL',' của hệ thống WebGL đã triển khai')
   break

# Update deployment tables, preserving their layout.
for t in doc.tables:
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for run in p.runs:
     run.text=run.text.replace('https://unimatch-web-20261016.vercel.app','https://unimatch-neu.vercel.app').replace('unimatch-web-20261016.vercel.app','unimatch-neu.vercel.app').replace('unimatch-web-20261016','unimatch').replace('unimatch-db-20261016','unimatch').replace('gentle-math-08514391','patient-tree-18935199').replace('https://github.com/namhv521/unimatch-deploy-20261016','https://github.com/namhv521/UniMatch').replace('https://github.com/namhv521/unimatch-uptime-20261016/actions','https://github.com/namhv521/UniMatch/actions').replace('unimatch-uptime-20261016','UniMatch/keepalive.yml').replace('unimatch-deploy-20261016','UniMatch').replace('Source deployment private','Mã nguồn công khai').replace('Lịch ping public','Lịch ping cùng repo')
# IEEE sources, including transparent use of former theses for structure only.
refs=next(p for p in doc.paragraphs if p.text=='PHỤ LỤC')
for text in [
 '[15] Three.js authors, “Three.js documentation,” [Trực tuyến]. Địa chỉ: https://threejs.org/docs/. [Truy cập: 16/10/2026].',
 '[16] N. H. Hải, “Ứng dụng mạng nơ-ron lượng tử trong phân loại cảm xúc văn bản tài chính,” Khóa luận tốt nghiệp, Trường Công nghệ, Đại học Kinh tế Quốc dân, 2026. Bản lưu trong data/1.pdfs, 92 trang.',
 '[17] N. Đ. Chung, “Thiết kế xây dựng phần mềm quản lý sinh viên,” Khóa luận tốt nghiệp, Trường Công nghệ, Đại học Kinh tế Quốc dân, 2025. Bản lưu trong data/1.pdfs, 74 trang.'
]:refs.insert_paragraph_before(text)
doc.add_heading('D Đối chiếu báo cáo khóa trước và quy định trình bày',level=2)
doc.add_paragraph('Báo cáo tham khảo cách tổ chức phần đầu, tổng quan, cơ sở lý thuyết, xây dựng hệ thống và thực nghiệm trong hai khóa luận [16], [17]. Báo cáo của Nguyễn Hoàng Hải được dùng để đối chiếu cách tách phương pháp với kết quả nghiên cứu; báo cáo của Nguyễn Đình Chung được dùng để đối chiếu phần phân tích và thiết kế phần mềm. Việc tham khảo chỉ phục vụ cấu trúc và trình bày, không dùng số liệu hay kết luận nghiên cứu của các tác giả làm kết quả UniMatch.')
doc.add_paragraph('Bản báo cáo chuyên đề thực tập về website trung tâm tiếng Anh của Lê Anh Minh trong cùng thư mục cũng được xem để phân biệt loại báo cáo. Quy định khóa luận của Khoa [14] là căn cứ ưu tiên khi các báo cáo cũ khác nhau về định dạng. Bản hiện hành dùng giấy A4, Times New Roman 13, giãn dòng 1,3, căn đều; lề trái 3,5 cm, các lề còn lại 2,5 cm. Tên chương cỡ 18, đậm và căn giữa; mục và tiểu mục cỡ 13, đậm. Nội dung chính được kiểm đếm từ Mở đầu đến hết Kết luận trên PDF đã xuất.')
doc.add_paragraph('Mục lục, danh mục bảng, hình, công thức và số trang được cập nhật bằng field của Word. Các tham khảo đánh số theo IEEE và liên kết bằng bookmark. Những dữ liệu chưa có nguồn chính thức vẫn để [Cần bổ sung]: tên giảng viên hướng dẫn, lớp, năm học/kỳ, chữ ký và kết quả Turnitin. Người viết cần hoàn tất những mục này trước khi ký và nộp; bản kỹ thuật không giả mạo xác nhận học vụ.')
# Preserve NEU requirements in sections and body styles.
for s in doc.sections:
 s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(3.5);s.right_margin=Cm(2.5);s.top_margin=Cm(2.5);s.bottom_margin=Cm(2.5)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(13);# Cover spacing is preserved from the NEU source template.
for paragraph_index,p in enumerate(doc.paragraphs):
 if paragraph_index < 51:continue
 if p.style.name=='Normal':p.paragraph_format.line_spacing=1.3;p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 if p.style.name=='Heading 1':p.style.font.size=Pt(18);p.style.font.name='Times New Roman';p.style.font.bold=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if p.style.name in ['Heading 2','Heading 3','Heading 4']:p.style.font.name='Times New Roman';p.style.font.size=Pt(13);p.style.font.bold=True
# Create automatic reference numbering and hyperlinkable numeric citations.
def field(run, code, cached):
 for tag,text in [('begin',None),('instruction',code),('separate',None),('text',cached),('end',None)]:
  el=OxmlElement('w:instrText' if tag=='instruction' else 'w:t' if tag=='text' else 'w:fldChar')
  if tag in ['instruction','text']:el.text=text
  else:el.set(qn('w:fldCharType'),tag)
  run._r.append(el)
references={}
for p in doc.paragraphs:
 m=re.match(r'^\[(\d+)\] ',p.text)
 if m:
  n=int(m[1]);text=p.text[m.end():];p.clear();p.add_run('[')
  start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(500+n));start.set(qn('w:name'),f'Reference_{n}');p._p.append(start)
  field(p.add_run(), ' SEQ Reference \\* ARABIC ',str(n))
  end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(500+n));p._p.append(end)
  p.add_run('] '+text);references[n]=True
for p in doc.paragraphs:
 if re.match(r'^\[\d+\] ',p.text) or p.style.name.startswith('TOC'):continue
 for run in list(p.runs):
  if not re.search(r'\[\d+\]',run.text):continue
  parts=re.split(r'(\[\d+\])',run.text);r=run._r;parent=r.getparent();index=parent.index(r)
  for part in parts:
   if not part:continue
   new=OxmlElement('w:r')
   if r.rPr is not None:new.append(deepcopy(r.rPr))
   parent.insert(index,new);index+=1
   from docx.text.run import Run
   wrapper=Run(new,p)
   if re.fullmatch(r'\[\d+\]',part) and int(part[1:-1]) in references:
    wrapper.add_text('[');field(wrapper,f' REF Reference_{part[1:-1]} \\h ',part[1:-1]);wrapper.add_text(']')
   else:wrapper.add_text(part)
  parent.remove(r)
settings=doc.settings.element
update=settings.find(qn('w:updateFields'))
if update is None:update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
output.parent.mkdir(exist_ok=True);doc.save(output)
print(output)
