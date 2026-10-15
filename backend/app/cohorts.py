from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

try:
    from lxml.etree import XMLSyntaxError
except ImportError:  # pragma: no cover - openpyxl can run without lxml
    XMLSyntaxError = ParseError

from src.baselines.hungarian import hungarian_assignment
from src.evaluation.metrics import load_variance, mean_compatibility, quota_violations
from src.features.feature_builder import FeatureBuilder

MAX_ZIP_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 2048
CAPACITY_POLICY = "balanced-divmod-selected-advisors-v1"
ALGORITHM = "hungarian-balanced-v1"
COHORT_TOKEN_PATTERN = r"(?u)(?:[A-Za-z]\+\+|[A-Za-z]#|\b\w+\b)"


class CohortError(ValueError):
    def __init__(self, message: str, *, status_code: int = 422, code: str = "invalid_cohort") -> None:
        super().__init__(message)
        self.status_code = status_code
        self.code = code


def _hash_json(value: object) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(payload).hexdigest()


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _skills(value: object) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for item in _text(value).split(";"):
        normalized = item.strip()
        folded = normalized.casefold()
        if normalized and folded not in seen:
            seen.add(folded)
            result.append(normalized)
    return result


class CohortService:
    def __init__(self, root: str | Path, matching_service) -> None:
        self.root = Path(root)
        self.matching_service = matching_service
        try:
            self.schema = json.loads(
                (self.root / "configs/student_upload_schema.json").read_text(encoding="utf-8")
            )
            survey_schema = json.loads(
                (self.root / "configs/student_survey_schema.json").read_text(encoding="utf-8")
            )
        except (OSError, KeyError, TypeError, json.JSONDecodeError) as exc:
            raise CohortError("Cấu hình nhập sinh viên không hợp lệ", status_code=503, code="cohort_unavailable") from exc
        self.schemas = {
            self.schema["sheet"]: self.schema,
            survey_schema["sheet"]: survey_schema,
        }
        self.columns = list(self.schema["columns"])
        self.column_keys = [column["key"] for column in self.columns]
        self.skill_keys = [column["key"] for column in self.columns if column["kind"] == "skills"]
        self.labels = {column["key"]: column["label"] for column in self.columns}

    def advisors(self) -> dict[str, object]:
        return {
            "advisors": [
                {
                    "id": row["advisor_id"],
                    "name": row["advisor_name"],
                    "primary_field": row.get("primary_field") or row.get("primary_field_inferred") or "",
                }
                for row in self.matching_service.advisors
            ]
        }

    def preview(self, content: bytes) -> dict[str, object]:
        parsed = self._parse(content)
        return {
            "student_count": len(parsed["students"]),
            "students": [
                ({
                    "student_id": row["student_id"],
                    "student_name": row["student_name"],
                    "thesis_title": row["thesis_title"],
                    "skills": row["skills"],
                } | ({"survey": row["survey"]} if "survey" in row else {}))
                for row in parsed["students"]
            ],
            "warnings": parsed["warnings"],
        }

    def allocate(self, content: bytes, advisor_ids: list[str]) -> dict[str, object]:
        selected = self._selected_advisors(advisor_ids)
        selected_advisor_ids = [row["advisor_id"] for row in selected]
        parsed = self._parse(content)
        students = parsed["students"]
        student_texts = [row["feature_text"] for row in students]
        all_advisor_texts = [self._advisor_text(row) for row in self.matching_service.advisors]
        train_indices = self.matching_service.split["indices"]["train"]
        train_texts = [
            self.matching_service.theses[index].get("feature_text")
            or self.matching_service.theses[index].get("thesis_title", "")
            for index in train_indices
        ]
        builder = FeatureBuilder()
        builder.vectorizer.set_params(token_pattern=COHORT_TOKEN_PATTERN)
        builder.fit(train_texts, all_advisor_texts)
        selected_texts = [self._advisor_text(row) for row in selected]
        scores, feature_manifest = builder.compatibility(student_texts, selected_texts)
        feature_manifest["token_pattern"] = COHORT_TOKEN_PATTERN

        quotient, remainder = divmod(len(students), len(selected))
        capacities = np.asarray(
            [quotient + (position < remainder) for position in range(len(selected))], dtype=int
        )
        assigned = hungarian_assignment(scores, capacities)
        loads = np.bincount(assigned, minlength=len(selected))
        warnings = list(parsed["warnings"])
        for position, row in enumerate(students):
            if float(np.max(scores[position])) <= 0:
                warnings.append(
                    f"Sinh viên {row['student_id']} chưa có kỹ năng trùng với hồ sơ giảng viên đã chọn"
                )

        advisors = [
            {
                "id": row["advisor_id"],
                "name": row["advisor_name"],
                "quota": int(capacities[position]),
                "assigned": int(loads[position]),
            }
            for position, row in enumerate(selected)
        ]
        assignments = [
            {
                "student_id": row["student_id"],
                "student_name": row["student_name"],
                "advisor_id": selected[int(assigned[position])]["advisor_id"],
                "advisor_name": selected[int(assigned[position])]["advisor_name"],
                "score": float(scores[position, assigned[position]]),
                "skills": row["skills"],
            }
            for position, row in enumerate(students)
        ]
        provenance = {
            "source_sha256": parsed["source_sha256"],
            "cohort_hash": parsed["cohort_hash"],
            "student_upload_schema_version": parsed["schema_version"],
            "source_schema_version": parsed["schema_version"],
            "source_sheet": parsed["sheet"],
            "advisor_snapshot": dict(self.matching_service.provenance),
            "feature_manifest": feature_manifest,
            "capacity_policy": CAPACITY_POLICY,
            "selected_advisor_ids": selected_advisor_ids,
        }
        return {
            "student_count": len(students),
            "advisor_count": len(selected),
            "source_sha256": parsed["source_sha256"],
            "cohort_hash": parsed["cohort_hash"],
            "capacity_policy": CAPACITY_POLICY,
            "algorithm": ALGORITHM,
            "selected_advisor_ids": selected_advisor_ids,
            "cohort": {"schema_version": parsed["schema_version"], "students": students},
            "advisors": advisors,
            "assignments": assignments,
            "metrics": {
                "mean_compatibility": mean_compatibility(scores, assigned),
                "load_variance": load_variance(assigned, len(selected)),
                "quota_violations": quota_violations(assigned, capacities),
                "advisor_loads": loads.tolist(),
            },
            "warnings": warnings,
            "provenance": provenance,
        }

    def _parse(self, content: bytes) -> dict[str, object]:
        max_bytes = int(self.schema["max_file_bytes"])
        if len(content) > max_bytes:
            raise CohortError("File Excel vượt quá giới hạn 5 MiB", status_code=413, code="file_too_large")
        if not content:
            raise CohortError("File Excel trống")
        self._validate_zip(content)
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, ParseError, XMLSyntaxError, InvalidFileException) as exc:
            raise CohortError("File XLSX không hợp lệ hoặc bị hỏng") from exc
        try:
            supported_sheets = [name for name in self.schemas if name in workbook.sheetnames]
            if not supported_sheets:
                raise CohortError("Thiếu sheet nhập liệu KhaoSat hoặc SinhVien")
            if len(supported_sheets) != 1:
                raise CohortError("File có nhiều sheet nhập liệu được hỗ trợ; chỉ giữ một sheet")
            sheet_name = supported_sheets[0]
            schema = self.schemas[sheet_name]
            columns = list(schema["columns"])
            column_keys = [column["key"] for column in columns]
            skill_keys = [column["key"] for column in columns if column["kind"] == "skills"]
            context_keys = [column["key"] for column in columns if column["kind"] == "context"]
            labels = {column["key"]: column["label"] for column in columns}
            expected_headers = [
                column["label"] if schema.get("header_mode") == "labels" else column["key"]
                for column in columns
            ]
            sheet = workbook[sheet_name]
            max_students = int(schema["max_students"])
            if sheet.max_row and sheet.max_row > max_students * 10 + 1:
                raise CohortError(f"Sheet {sheet_name} khai báo quá nhiều dòng để xử lý an toàn")
            header_cells = next(sheet.iter_rows(min_row=int(schema["header_row"]), max_row=int(schema["header_row"])))
            headers = [_text(cell.value) for cell in header_cells]
            if any(cell.data_type == "f" for cell in header_cells):
                raise CohortError("Tiêu đề cột không được chứa công thức")
            if len(headers) != len(set(headers)):
                raise CohortError("Tiêu đề cột bị trùng")
            missing = [header for header in expected_headers if header not in headers]
            unknown = [header for header in headers if header not in expected_headers]
            if missing or unknown:
                detail = ", ".join([*(f"thiếu {key}" for key in missing), *(f"không hỗ trợ {key}" for key in unknown)])
                raise CohortError(f"Tiêu đề cột không đúng: {detail}")
            column_index = {
                key: headers.index(header)
                for key, header in zip(column_keys, expected_headers, strict=True)
            }
            students: list[dict[str, object]] = []
            warnings: list[str] = []
            seen_ids: set[str] = set()
            blank_rows = 0
            for excel_row, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
                values = [cell.value for cell in cells]
                if not any(_text(value) for value in values):
                    blank_rows += 1
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise CohortError(f"Dòng {excel_row}: không được dùng công thức")
                raw = {
                    key: values[index] if index < len(values) else None
                    for key, index in column_index.items()
                }
                student_id = _text(raw["student_id"])
                student_name = _text(raw["student_name"])
                if not student_id:
                    raise CohortError(f"Dòng {excel_row}: thiếu {labels['student_id']}")
                if not student_name:
                    raise CohortError(f"Dòng {excel_row}: thiếu {labels['student_name']}")
                if student_id in seen_ids:
                    raise CohortError(f"Dòng {excel_row}: Mã sinh viên bị trùng: {student_id}")
                seen_ids.add(student_id)
                skill_fields = {key: _skills(raw[key]) for key in skill_keys}
                flattened: list[str] = []
                folded: set[str] = set()
                for key in skill_keys:
                    for skill in skill_fields[key]:
                        if skill.casefold() not in folded:
                            folded.add(skill.casefold())
                            flattened.append(skill)
                if not flattened:
                    raise CohortError(f"Dòng {excel_row}: cần ít nhất một kỹ năng")
                thesis_title = _text(raw["thesis_title"])
                abstract = _text(raw["abstract"])
                student = {
                    "student_id": student_id,
                    "student_name": student_name,
                    "thesis_title": thesis_title,
                    "abstract": abstract,
                    **skill_fields,
                    "skills": flattened,
                }
                student["feature_text"] = " ".join(
                    value for value in [thesis_title, abstract, *flattened] if value
                )
                if schema["version"] == "2.0":
                    declared = {skill.casefold(): skill for skill in flattened}
                    levels = self._skill_levels(raw.get("skill_levels"), declared, excel_row)
                    student["survey"] = {
                        "schema_version": schema["version"],
                        "raw": {key: _text(raw.get(key)) for key in column_keys},
                        "context": {key: _text(raw.get(key)) for key in context_keys},
                        "skill_levels": levels,
                        "supplemental_skills": {
                            key: skill_fields[key]
                            for key in ("mobile_frameworks", "dev_tools")
                        },
                    }
                students.append(student)
                if len(students) > max_students:
                    raise CohortError("File chứa tối đa 1000 sinh viên")
            if not students:
                raise CohortError("File không có sinh viên hợp lệ")
            if blank_rows:
                warnings.append(f"{blank_rows} dòng trống được bỏ qua")
        except CohortError:
            raise
        except (OSError, ValueError, KeyError, ParseError, XMLSyntaxError, InvalidFileException) as exc:
            raise CohortError("File XLSX không hợp lệ hoặc bị hỏng") from exc
        finally:
            workbook.close()
        return {
            "students": students,
            "warnings": warnings,
            "source_sha256": hashlib.sha256(content).hexdigest(),
            "cohort_hash": _hash_json(students),
            "schema_version": schema["version"],
            "sheet": sheet_name,
        }

    @staticmethod
    def _skill_levels(value: object, declared: dict[str, str], excel_row: int) -> dict[str, int]:
        levels: dict[str, int] = {}
        seen: set[str] = set()
        for raw_item in _text(value).split(";"):
            item = raw_item.strip()
            if not item:
                continue
            parts = [part.strip() for part in item.split("|")]
            if len(parts) != 2 or not all(parts):
                raise CohortError(f"Dòng {excel_row}: mức kỹ năng phải theo định dạng Kỹ năng|Mức")
            skill, level = parts
            folded = skill.casefold()
            if folded in seen:
                raise CohortError(f"Dòng {excel_row}: mức kỹ năng bị trùng: {skill}")
            seen.add(folded)
            if folded not in declared:
                raise CohortError(f"Dòng {excel_row}: kỹ năng trong mức độ chưa được khai báo: {skill}")
            if level not in {"1", "2", "3"}:
                raise CohortError(f"Dòng {excel_row}: mức kỹ năng chỉ nhận 1, 2 hoặc 3")
            levels[declared[folded]] = int(level)
        return levels

    @staticmethod
    def _validate_zip(content: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_ZIP_ENTRIES:
                    raise CohortError("File XLSX có quá nhiều thành phần")
                if any(entry.flag_bits & 0x1 for entry in entries):
                    raise CohortError("File XLSX mã hóa không được hỗ trợ")
                if sum(entry.file_size for entry in entries) > MAX_ZIP_EXPANDED_BYTES:
                    raise CohortError("File XLSX giải nén vượt quá giới hạn an toàn")
        except zipfile.BadZipFile as exc:
            raise CohortError("File XLSX không hợp lệ hoặc bị hỏng") from exc

    def _selected_advisors(self, advisor_ids: list[str]) -> list[dict[str, object]]:
        if not advisor_ids:
            raise CohortError("Chọn ít nhất một giảng viên")
        if len(advisor_ids) != len(set(advisor_ids)):
            raise CohortError("Giảng viên được chọn bị trùng")
        by_id = {row["advisor_id"]: row for row in self.matching_service.advisors}
        unknown = [advisor_id for advisor_id in advisor_ids if advisor_id not in by_id]
        if unknown:
            raise CohortError(f"Giảng viên không tồn tại: {', '.join(unknown)}")
        selected = set(advisor_ids)
        return [row for row in self.matching_service.advisors if row["advisor_id"] in selected]

    @staticmethod
    def _advisor_text(row: dict[str, object]) -> str:
        return " ".join(
            str(row.get(key) or "")
            for key in ("profile_text", "skill_text", "skill_text_inferred", "primary_field")
        ).strip()
