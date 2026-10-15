from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from backend.app.cohorts import CohortError, CohortService
from src.serving import MatchingService

ROOT = Path(__file__).parents[2]
SCHEMA = json.loads((ROOT / "configs/student_upload_schema.json").read_text(encoding="utf-8"))
HEADERS = [column["key"] for column in SCHEMA["columns"]]
SURVEY_SCHEMA = json.loads((ROOT / "configs/student_survey_schema.json").read_text(encoding="utf-8"))
SURVEY_LABELS = [column["label"] for column in SURVEY_SCHEMA["columns"]]


def workbook_bytes(rows: list[dict[str, object]], *, headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SinhVien"
    sheet.append(headers or HEADERS)
    for row in rows:
        sheet.append([row.get(key) for key in (headers or HEADERS)])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def survey_workbook_bytes(rows: list[dict[str, object]], *, include_legacy_sheet: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KhaoSat"
    sheet.append(SURVEY_LABELS)
    keys = [column["key"] for column in SURVEY_SCHEMA["columns"]]
    for row in rows:
        sheet.append([row.get(key) for key in keys])
    if include_legacy_sheet:
        legacy = workbook.create_sheet("SinhVien")
        legacy.append(HEADERS)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def workbook_with_declared_max_row(content: bytes, max_row: int) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(content))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b'ref="A1:Q2"', f'ref="A1:Q{max_row}"'.encode())
            target.writestr(entry, payload)
    return output.getvalue()


def workbook_with_corrupt_sheet_xml(content: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(content))
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            payload = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                payload = b"<worksheet><broken>"
            target.writestr(entry, payload)
    return output.getvalue()


def student(position: int, *, skill: str = "Python; Machine Learning") -> dict[str, object]:
    return {
        "student_id": f"SV{position:03}",
        "student_name": f"Sinh viên {position}",
        "thesis_title": f"Đề tài {position}",
        "abstract": "Ứng dụng phân tích dữ liệu",
        "app_languages": skill,
    }


def survey_student(**overrides: object) -> dict[str, object]:
    return {
        "student_id": "KS001",
        "student_name": "Sinh viên khảo sát",
        "class_name": "K66 CNTT",
        "major": "Công nghệ thông tin",
        "app_languages": "Python",
        "skill_levels": "Python|2",
        "project_experience": "Làm API môn học",
        "project_links": "https://example.test/student",
        "preferred_fields": "Thị giác máy tính",
        "learning_goals": "Muốn học PyTorch",
        "thesis_title": "",
        "abstract": "",
        **overrides,
    }


@pytest.fixture(scope="module")
def cohort_service() -> CohortService:
    return CohortService(ROOT, MatchingService(ROOT))


def test_preview_normalizes_semicolon_skills_and_keeps_real_student_identity(cohort_service) -> None:
    preview = cohort_service.preview(workbook_bytes([student(1)]))

    assert preview["student_count"] == 1
    assert preview["students"] == [{
        "student_id": "SV001",
        "student_name": "Sinh viên 1",
        "thesis_title": "Đề tài 1",
        "skills": ["Python", "Machine Learning"],
    }]


def test_survey_accepts_student_without_topic_and_preserves_context_and_levels(cohort_service) -> None:
    preview = cohort_service.preview(survey_workbook_bytes([survey_student()]))

    assert preview["student_count"] == 1
    row = preview["students"][0]
    assert row["thesis_title"] == ""
    assert row["skills"] == ["Python"]
    assert row["survey"]["context"] == {
        "class_name": "K66 CNTT",
        "major": "Công nghệ thông tin",
        "project_experience": "Làm API môn học",
        "project_links": "https://example.test/student",
        "preferred_fields": "Thị giác máy tính",
        "learning_goals": "Muốn học PyTorch",
    }
    assert row["survey"]["skill_levels"] == {"Python": 2}


def test_survey_context_is_retained_but_excluded_from_matching_text(cohort_service) -> None:
    marker = "CONTEXT-MUST-NOT-BE-SCORED-XYZ"
    parsed = cohort_service._parse(survey_workbook_bytes([survey_student(learning_goals=marker)]))

    row = parsed["students"][0]
    assert row["survey"]["context"]["learning_goals"] == marker
    assert marker not in row["feature_text"]


@pytest.mark.parametrize(
    ("levels", "message"),
    [
        ("Python|4", "chỉ nhận 1, 2 hoặc 3"),
        ("Rust|2", "chưa được khai báo"),
        ("Python|2; python|3", "bị trùng"),
        ("Python|2; Python|2", "bị trùng"),
        ("Python-2", "định dạng Kỹ năng|Mức"),
    ],
)
def test_survey_rejects_invalid_skill_levels(cohort_service, levels, message) -> None:
    with pytest.raises(CohortError, match=message):
        cohort_service.preview(survey_workbook_bytes([survey_student(skill_levels=levels)]))


def test_survey_supplemental_skills_are_matched_but_kept_separate(cohort_service) -> None:
    parsed = cohort_service._parse(
        survey_workbook_bytes([
            survey_student(
                app_languages="Python",
                mobile_frameworks="Flutter",
                dev_tools="Docker; Git",
                skill_levels="Python|2; Flutter|1; Docker|3",
            )
        ])
    )

    row = parsed["students"][0]
    assert row["skills"] == ["Python", "Flutter", "Docker", "Git"]
    assert row["mobile_frameworks"] == ["Flutter"]
    assert row["dev_tools"] == ["Docker", "Git"]
    assert row["survey"]["supplemental_skills"] == {
        "mobile_frameworks": ["Flutter"],
        "dev_tools": ["Docker", "Git"],
    }
    assert all(skill in row["feature_text"] for skill in ("Python", "Flutter", "Docker", "Git"))


def test_survey_allocation_records_v2_schema_and_source_sheet(cohort_service) -> None:
    selected = [advisor["id"] for advisor in cohort_service.advisors()["advisors"][:2]]

    result = cohort_service.allocate(
        survey_workbook_bytes([survey_student(thesis_title="", abstract="")]), selected
    )

    assert result["cohort"]["schema_version"] == "2.0"
    assert result["provenance"]["source_schema_version"] == "2.0"
    assert result["provenance"]["source_sheet"] == "KhaoSat"
    assert result["cohort"]["students"][0]["survey"]["raw"]["thesis_title"] == ""


def test_workbook_with_both_supported_input_sheets_is_rejected(cohort_service) -> None:
    with pytest.raises(CohortError, match="nhiều sheet nhập liệu"):
        cohort_service.preview(
            survey_workbook_bytes([survey_student()], include_legacy_sheet=True)
        )


@pytest.mark.parametrize(
    ("rows", "headers", "message"),
    [
        ([student(1), student(1)], None, "Mã sinh viên bị trùng"),
        ([{**student(1), "student_name": ""}], None, "Dòng 2: thiếu Họ và tên"),
        ([{**student(1), "app_languages": ""}], None, "Dòng 2: cần ít nhất một kỹ năng"),
        ([student(1)], [*HEADERS[:-1], HEADERS[-2]], "Tiêu đề cột bị trùng"),
    ],
)
def test_preview_rejects_invalid_rows_with_useful_vietnamese_errors(
    cohort_service, rows, headers, message
) -> None:
    with pytest.raises(CohortError, match=message):
        cohort_service.preview(workbook_bytes(rows, headers=headers))


def test_preview_rejects_formulas_and_file_size_boundary(cohort_service) -> None:
    with pytest.raises(CohortError, match="Dòng 2.*công thức"):
        cohort_service.preview(workbook_bytes([{**student(1), "app_languages": "=1+1"}]))

    with pytest.raises(CohortError, match="5 MiB"):
        cohort_service.preview(b"x" * (5 * 1024 * 1024 + 1))


def test_preview_rejects_zip_expansion_bomb(cohort_service) -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/oversized.xml", b"0" * (50 * 1024 * 1024 + 1))

    with pytest.raises(CohortError, match="giải nén vượt quá giới hạn an toàn"):
        cohort_service.preview(output.getvalue())


def test_preview_rejects_more_than_1000_students(cohort_service) -> None:
    with pytest.raises(CohortError, match="tối đa 1000 sinh viên"):
        cohort_service.preview(workbook_bytes([student(index) for index in range(1, 1002)]))


def test_preview_aggregates_styled_empty_template_rows(cohort_service) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SinhVien"
    sheet.append(HEADERS)
    sheet.append([student(1).get(key) for key in HEADERS])
    for row in range(3, 7):
        for column in range(1, len(HEADERS) + 1):
            sheet.cell(row=row, column=column).number_format = "@"
    output = io.BytesIO()
    workbook.save(output)

    preview = cohort_service.preview(output.getvalue())

    assert preview["warnings"] == ["4 dòng trống được bỏ qua"]


def test_preview_rejects_worksheet_declaring_excessive_scan_range(cohort_service) -> None:
    content = workbook_with_declared_max_row(workbook_bytes([student(1)]), 1_048_576)

    with pytest.raises(CohortError, match="khai báo quá nhiều dòng"):
        cohort_service.preview(content)


def test_preview_rejects_corrupt_sheet_xml_as_validation_error(cohort_service) -> None:
    content = workbook_with_corrupt_sheet_xml(workbook_bytes([student(1)]))

    with pytest.raises(CohortError, match="không hợp lệ hoặc bị hỏng"):
        cohort_service.preview(content)


def test_balanced_hungarian_fills_exact_divmod_quotas_for_38_students(cohort_service) -> None:
    selected = [advisor["id"] for advisor in cohort_service.advisors()["advisors"][:4]]

    result = cohort_service.allocate(
        workbook_bytes([student(index) for index in range(1, 39)]), selected
    )

    assert result["student_count"] == 38
    assert result["advisor_count"] == 4
    assert [advisor["quota"] for advisor in result["advisors"]] == [10, 10, 9, 9]
    assert [advisor["assigned"] for advisor in result["advisors"]] == [10, 10, 9, 9]
    assert {row["advisor_id"] for row in result["assignments"]} <= set(selected)
    assert result["algorithm"] == "hungarian-balanced-v1"
    assert result["metrics"]["quota_violations"] == 0


def test_reversed_advisor_query_uses_same_canonical_order_and_assignments(cohort_service) -> None:
    selected = [advisor["id"] for advisor in cohort_service.advisors()["advisors"][:4]]
    content = workbook_bytes([student(index) for index in range(1, 39)])

    forward = cohort_service.allocate(content, selected)
    reversed_result = cohort_service.allocate(content, list(reversed(selected)))

    assert [row["id"] for row in reversed_result["advisors"]] == selected
    assert reversed_result["selected_advisor_ids"] == selected
    assert reversed_result["provenance"]["selected_advisor_ids"] == selected
    assert reversed_result["assignments"] == forward["assignments"]


def test_cohort_tokenizer_preserves_cpp_csharp_and_single_letter_r_skills() -> None:
    advisors = [
        {"advisor_id": "GV-CPP", "advisor_name": "GV C++", "skill_text": "C++"},
        {"advisor_id": "GV-CSHARP", "advisor_name": "GV C#", "skill_text": "C#"},
        {"advisor_id": "GV-R", "advisor_name": "GV R", "skill_text": "R"},
    ]
    matching = SimpleNamespace(
        advisors=advisors,
        theses=[{"feature_text": "baseline corpus"}],
        split={"indices": {"train": [0]}},
        provenance={"advisor_order_sha256": "synthetic-token-regression"},
    )
    service = CohortService(ROOT, matching)
    rows = [
        {**student(1, skill="C++"), "thesis_title": "", "abstract": ""},
        {**student(2, skill="C#"), "thesis_title": "", "abstract": ""},
        {**student(3, skill="R"), "thesis_title": "", "abstract": ""},
    ]

    result = service.allocate(workbook_bytes(rows), [row["advisor_id"] for row in advisors])

    assert [row["advisor_id"] for row in result["assignments"]] == ["GV-CPP", "GV-CSHARP", "GV-R"]
    assert all(row["score"] > 0 for row in result["assignments"])
    assert result["provenance"]["feature_manifest"]["token_pattern"] == (
        r"(?u)(?:[A-Za-z]\+\+|[A-Za-z]#|\b\w+\b)"
    )


def test_allocate_rejects_empty_duplicate_and_unknown_advisor_ids(cohort_service) -> None:
    content = workbook_bytes([student(1)])
    advisor_id = cohort_service.advisors()["advisors"][0]["id"]

    with pytest.raises(CohortError, match="Chọn ít nhất một giảng viên"):
        cohort_service.allocate(content, [])
    with pytest.raises(CohortError, match="Giảng viên được chọn bị trùng"):
        cohort_service.allocate(content, [advisor_id, advisor_id])
    with pytest.raises(CohortError, match="Giảng viên không tồn tại"):
        cohort_service.allocate(content, ["unknown-advisor"])


def test_more_advisors_than_students_assigns_zero_quotas_and_warns_zero_overlap(cohort_service) -> None:
    selected = [advisor["id"] for advisor in cohort_service.advisors()["advisors"][:4]]
    row = {**student(1, skill="unknown-skill-token-zzzz"), "thesis_title": "", "abstract": ""}

    result = cohort_service.allocate(
        workbook_bytes([row]), selected
    )

    assert [advisor["quota"] for advisor in result["advisors"]] == [1, 0, 0, 0]
    assert [advisor["assigned"] for advisor in result["advisors"]] == [1, 0, 0, 0]
    assert any("chưa có kỹ năng trùng" in warning for warning in result["warnings"])

